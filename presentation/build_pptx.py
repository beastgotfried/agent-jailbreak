from pathlib import Path
import csv
import json
import re
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'Veyra-IBM-Z-Datathon.pptx'
ASSETS=ROOT/'ppt-assets';ASSETS.mkdir(exist_ok=True)
metrics=json.loads((ROOT/'evidence/metrics.json').read_text())
content=(ROOT/'Veyra-Pitch-Content.txt').read_text()
parts=re.split(r'(?m)^## Slide (\d+) — ',content)
notes={int(parts[i]):parts[i+1].split('\n---\n')[0] for i in range(1,len(parts),2)}
refs=content.split('## Appendix G — Source and Evidence References',1)[1].split('## Appendix H',1)[0]

BLUE='0F62FE';NAVY='001141';INK='161616';TEAL='007D79';GRAY='525252';LIGHT='EDF5FF';LINE='DDE1E6';WHITE='FFFFFF'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='Veyra — IBM Z Datathon'
prs.core_properties.subject='Native LinuxONE guardrail integrity and verified recovery'
prs.core_properties.author='Veyra'
prs.core_properties.keywords='IBM Z, LinuxONE, s390x, guardrail, verified recovery'
prs.core_properties.comments='Frozen release 8a86ce0. Detailed evidence, scope and sources appear in speaker notes.'

def rgb(s):return RGBColor.from_string(s)
def rect(slide,x,y,w,h,fill=WHITE,line=None,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill)
    if line:shape.line.color.rgb=rgb(line);shape.line.width=Pt(1)
    else:shape.line.fill.background()
    if radius:shape.adjustments[0]=.1
    return shape

def text(slide,x,y,w,h,value,size=22,color=INK,bold=False,align=PP_ALIGN.LEFT,font='Arial',valign=MSO_ANCHOR.TOP):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.vertical_anchor=valign
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align
        p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        p.space_before=Pt(0);p.space_after=Pt(size*.2);p.line_spacing=1.08
    return shape

def arrow(slide,x1,y1,x2,y2,color=BLUE,width=1.8):
    shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    shape.line.color.rgb=rgb(color);shape.line.width=Pt(width)
    el=OxmlElement('a:tailEnd');el.set('type','triangle');shape.line._get_or_add_ln().append(el)
    return shape

def footer(slide,n,source,dark=False):
    color='A6C8FF' if dark else '6F6F6F'
    text(slide,.62,7.08,11.7,.22,source,10.5,color)
    text(slide,12.4,7.06,.3,.22,f'{n:02d}',10.5,color,align=PP_ALIGN.RIGHT,font='Courier New')

def base(n,kicker,title,sub=''):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(WHITE)
    rect(slide,0,0,.13,7.5,BLUE)
    text(slide,.62,.32,11.8,.25,kicker.upper(),12,BLUE,True,font='Arial')
    text(slide,.62,.85,12.05,.68,title,34,INK,True)
    if sub:text(slide,.65,1.55,12,.42,sub,16,GRAY)
    slide.notes_slide.notes_text_frame.text=notes[n]+'\n\nSOURCE REFERENCE INDEX\n'+refs
    return slide

def node(s,x,y,w,h,title,detail,fill=WHITE,color=BLUE):
    rect(s,x,y,w,h,fill,'A6C8FF',True)
    text(s,x+.13,y+.12,w-.26,.3,title,18,INK,True,PP_ALIGN.CENTER)
    text(s,x+.13,y+.49,w-.26,h-.54,detail,14,GRAY,align=PP_ALIGN.CENTER)

# 01: title and native execution promise.
s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(NAVY)
rect(s,.62,.5,.5,.06,'78A9FF')
text(s,.62,.75,8,.3,'IBM Z DATATHON  /  NATIVE LINUXONE',13,'A6C8FF',True)
text(s,.62,1.55,7.7,2.1,'Integrity\nWatch',72,WHITE,True)
text(s,.68,4.15,7.2,1.45,'Observe change.\nVerify recovery. Keep control.',28,'D0E2FF')
text(s,.68,6.15,7.6,.55,'An evidence-backed control loop around AI guardrails.',19,WHITE)
for i,(a,b) in enumerate([('01','Native s390x CPU execution'),('02','Matched runtime integrity checks'),('03','Approved, verified recovery')]):
    y=2.15+i*1.1
    text(s,8.65,y,.5,.3,a,14,'78A9FF',True,font='Courier New')
    text(s,9.3,y,3.25,.7,b,22,WHITE,True)
    rect(s,8.65,y+.82,3.9,.012,'244A83')
footer(s,1,'Frozen offline research release 8a86ce0 · Source and measured scope in speaker notes',True)
s.notes_slide.notes_text_frame.text=notes[1]+'\n\nSOURCE REFERENCE INDEX\n'+refs

# 02: actual recorded score transformation, without an invented request label.
with (ROOT/'tables/scale-timeline.csv').open() as f:rows=list(csv.DictReader(f))
threshold=.20832091569900513
example=next(r for r in rows if 128<=int(r['index'])<175 and float(r['healthy_guardrail_score'])>=threshold and float(r['A_guardrail_score'])<threshold)
s=base(2,'The problem','Valid scores can hide a changed control.','A changed output path can change decisions while the model still returns valid numbers.')
text(s,.68,2.35,5.05,2.0,'Same model.\nSame request.\nChanged runtime output.',30,INK,True)
text(s,.68,4.85,5.0,1.12,'One score cannot tell us whether traffic changes—or the control changes.',22,GRAY)
rect(s,6.15,2.25,6.45,3.92,LIGHT,None,True)
text(s,6.48,2.58,5.8,.3,'ACTUAL NATIVE REPLAY EXAMPLE',13,BLUE,True)
text(s,6.5,3.12,2.35,.3,'Approved baseline',16,GRAY)
text(s,9.9,3.12,2.35,.3,'After ×0.25 fault',16,GRAY)
text(s,6.5,3.6,2.4,.8,f'{float(example["healthy_guardrail_score"]):.4f}',43,INK,True)
text(s,9.9,3.6,2.35,.8,f'{float(example["A_guardrail_score"]):.4f}',43,BLUE,True)
arrow(s,8.9,4.05,9.65,4.05)
text(s,6.5,4.72,5.65,.45,'Fixed decision threshold: 0.2083',20,INK,True)
text(s,6.5,5.38,5.7,.4,f'Recorded scale / seed 0 · Request index {example["index"]}',14,GRAY)
footer(s,2,'Evidence: FaultBackend.score · Hash-verified scale-0 event trace · Example shows a decision change, not a new safety label')

# 03: roles and authority.
s=base(3,'The solution','Observe. Verify. Respond.','Separate the request decision, runtime evidence, and response authority.')
cols=[(.65,'01','REQUEST DECISION','Guardrail + canary','The guardrail gates the request.\nThe canary supplies a second score view.'),
      (4.87,'02','RUNTIME EVIDENCE','Matched probes','Check both score vectors against a fixed, approved baseline.'),
      (9.08,'03','RESPONSE AUTHORITY','Verified recovery','Execute the candidate. Verify it.\nApply only with approval.')]
for x,num,kick,title,body in cols:
    rect(s,x,2.3,3.61,.045,BLUE)
    text(s,x,2.65,3.6,.3,num,15,BLUE,True,font='Courier New')
    text(s,x,3.18,3.6,.65,title,27,INK,True)
    text(s,x,4.05,3.6,1.45,body,21,GRAY)
    text(s,x,5.7,3.6,.35,kick,12,BLUE,True)
rect(s,.65,6.35,12.03,.4,LIGHT)
text(s,.86,6.42,11.6,.25,'A traffic warning is evidence—not authorization to replace the model.',16,BLUE,True)
footer(s,3,'Evidence: SupervisedRuntime · RuntimeWatch · RecoveryManager · ProfiledTrafficWatch')

# 04: editable implemented architecture.
s=base(4,'Implemented architecture','One native execution boundary.')
rect(s,.48,1.72,12.38,4.98,LIGHT,'A6C8FF',True)
text(s,.76,1.9,11.8,.28,'IBM LINUXONE  /  LINUX s390x  /  CPU-ONLY EXECUTION',15,BLUE,True)
xs=[.78,3.91,7.04,10.17]
for x,title,detail in zip(xs,['Request stream','Guardrail + canary','Frozen policy','Agent stub'],['Ordered Python replay','Actual inference','Guardrail decision','Allow / refuse / hold']):node(s,x,2.52,2.39,1.02,title,detail)
for x in xs[:-1]:arrow(s,x+2.43,3.02,x+3.04,3.02)
node(s,1.03,4.0,3.2,1.05,'RuntimeWatch','8 probes / every 16 requests')
node(s,5.06,4.0,3.2,1.05,'RecoveryManager','Affine candidate / trusted reload')
node(s,9.09,4.0,3.2,1.05,'Verify + activate','Fresh paired check + approval')
arrow(s,4.27,4.52,4.96,4.52,TEAL);arrow(s,8.3,4.52,8.99,4.52,TEAL)
arrow(s,2.63,3.98,2.63,3.79,BLUE)
arrow(s,2.63,3.79,5.12,3.79,BLUE)
arrow(s,5.12,3.79,5.12,3.58,BLUE)
text(s,3.02,3.56,2.55,.23,'Matched probe execution',11.5,BLUE,align=PP_ALIGN.CENTER)
arrow(s,10.7,3.98,10.7,3.58,TEAL)
text(s,10.7,3.67,1.75,.22,'Or explicit hold',11.5,TEAL)
rect(s,.78,5.52,3.72,.7,WHITE)
text(s,.91,5.65,3.46,.4,'Frozen models, policy\nand disjoint labelled probes',14,GRAY)
rect(s,4.81,5.52,3.72,.7,WHITE)
text(s,4.94,5.65,3.46,.4,'Separate KS traffic study:\nWarnings have no action authority',14,GRAY)
rect(s,8.84,5.52,3.72,.7,WHITE)
text(s,8.97,5.65,3.46,.4,'Local JSON / JSONL evidence\nand SHA-256 manifests',14,GRAY)
text(s,.78,6.37,11.8,.23,'Current interfaces: CLI + Python API. Enterprise web console, REST intake, and durable services remain roadmap items.',11.7,GRAY)
footer(s,4,'Evidence: src/monitor/runtime.py · src/response/supervised.py · src/response/recovery.py · Full Mermaid diagram in the content kit')

# 05: dataset and CPU models; full source/licence table remains in notes.
s=base(5,'Data + model development','Controlled data. Lightweight models.','100,000 synthetic requests · 100 template families · 91 reviewed semantic groups')
left=.68
for value,fill,label in [(70,BLUE,'70,000 train'),(10,TEAL,'10,000\nvalidation'),(20,'A6C8FF','20,000 test')]:
    width=12*value/100;rect(s,left,2.15,width,.77,fill)
    text(s,left+.1,2.31,width-.2,.44,label,20 if value!=10 else 13,INK if value==20 else WHITE,True,PP_ALIGN.CENTER)
    left+=width
text(s,.7,3.12,12,.35,'Related cases stay together. Only text enters the classifiers.',19,GRAY)
for x,title,lines in [(.7,'Guardrail',['Word + character TF-IDF','Logistic regression · C=4','Frozen threshold · 0.2083']),
                     (4.98,'Canary',['Character TF-IDF','Logistic regression · C=1','Frozen threshold · 0.2771'])]:
    rect(s,x,3.77,3.85,1.76,LIGHT,None,True)
    text(s,x+.18,3.95,3.48,.34,title,23,INK,True)
    text(s,x+.18,4.48,3.48,.85,'\n'.join(lines),16,GRAY)
rect(s,9.26,3.77,3.39,1.76,'F4F4F4',None,True)
text(s,9.44,3.95,3.03,.34,'Healthy test · 20k',20,INK,True)
text(s,9.44,4.48,3.03,.85,'88.67% accuracy\n99.96% harmful recall\n22.63% safe FPR',16,GRAY)
text(s,.7,5.92,12,.4,'External regression: XSTest (300) + human ToxicChat (2,853).',18,INK,True)
text(s,.7,6.38,12,.43,'Earlier sources: WAInjectBench, JBB + mirror, StrongREJECT, HarmBench, Dolly. OASST1: acquisition only. Full roles and licenses in notes.',12.5,GRAY)
footer(s,5,'Evidence: data-manifest.json · model-cards.json · benchmark-report.json · Variants are not independent human requests')

# 06: distinct detector paths and researched methods.
s=base(6,'Detection methodology','Two signals. Different authority.','Traffic change, runtime change, and labelled safety performance remain separate.')
text(s,.72,2.2,11.9,.3,'HISTORICAL TRAFFIC-STUDY PATH',13,GRAY,True)
for x,title,detail in [(1.0,'Signed score gap','Guardrail score − canary score'),(5.0,'Empirical KS','Window 64 / stride 16'),(9.0,'Traffic warning','No response authority')]:node(s,x,2.67,3.3,1.0,title,detail,fill='F4F4F4')
arrow(s,4.34,3.17,4.91,3.17,GRAY);arrow(s,8.34,3.17,8.91,3.17,GRAY)
text(s,.72,4.02,11.9,.3,'SELECTED RELEASE: RUNTIME PATH',13,BLUE,True)
for x,title,detail in [(1.0,'8 fixed probes','Every 16 request positions'),(5.0,'Compare both models','Tolerance: 1e-8'),(9.0,'Verify response','Actual candidate execution')]:node(s,x,4.5,3.3,1.0,title,detail,fill=LIGHT)
arrow(s,4.34,5.0,4.91,5.0);arrow(s,8.34,5.0,8.91,5.0)
text(s,.76,6.0,11.8,.35,'Project contribution: verified recovery, explicit holds, repeated incidents, and native reproducibility.',18,INK,True)
text(s,.76,6.5,11.8,.3,'Cheap Canaries is prior art. Scan martingales and MMD are reference methods, not implemented in this release.',13,GRAY)
footer(s,6,'Evidence: src/monitor/disagreement.py + runtime.py · Cheap Canaries source snapshot 90726d9 · No head-to-head performance claim')

# 07: actual timeline figure, sized for the slide.
plt.rcParams.update({'font.family':'Arial','font.size':12,'svg.fonttype':'none'})
fig,ax=plt.subplots(figsize=(12,3.55));fig.subplots_adjust(left=.07,right=.985,bottom=.2,top=.89)
x=[int(r['index']) for r in rows]
a=[float(r['A_absolute_baseline_error']) for r in rows];b=[float(r['B_absolute_baseline_error']) for r in rows]
ax.axvspan(128,175,color='#edf5ff');ax.plot(x,a,color='#8d8d8d',lw=1.2,label='Guardrail only');ax.plot(x,b,color='#0f62fe',lw=2,label='Monitored + verified response')
for i,col in [(128,'#525252'),(143,'#007d79'),(175,'#0f62fe')]:ax.axvline(i,color=col,ls='--',lw=1)
ax.set_xlim(0,511);ax.set_ylim(-.025,.84);ax.set_xlabel('Request index (zero-based)');ax.set_ylabel('Absolute score error');ax.grid(axis='y',color='#e0e0e0',lw=.5);ax.spines[['top','right']].set_visible(False);ax.legend(frameon=False,loc='upper right',ncol=2,fontsize=11)
fig.savefig(ASSETS/'runtime-trace.png',dpi=200);plt.close(fig)
s=base(7,'Product demonstration','From fault to verified recovery.','Actual native scale / seed-0 trace—not a product dashboard or a simulated screenshot.')
for x,num,title,col in [(.76,'128','FAULT STARTS',GRAY),(4.98,'143','RUNTIME CONFIRMS',TEAL),(9.2,'175','VERIFIED CORRECTION',BLUE)]:
    text(s,x,2.09,1,.52,num,32,col,True);text(s,x+1.1,2.22,2.75,.3,title,13,col,True)
s.shapes.add_picture(str(ASSETS/'runtime-trace.png'),Inches(.55),Inches(2.95),width=Inches(12.2),height=Inches(3.6))
text(s,.76,6.52,11.8,.35,'1,000-request endpoint: guardrail error 0.0 · canary error 0.0 · decision mismatches 0',16,BLUE,True)
footer(s,7,'Evidence: hash-verified scale-0 events and endpoint · 15-request confirmation delay + 32-position response delay')

# 08: complete denominators and trade-offs.
s=base(8,'Measured results','Measure every request—not only the endpoint.','60 paired episodes · 30,720 observations · Controlled synthetic fault suite')
text(s,.72,2.15,5.6,.36,'CORRECT AUTOMATIC DECISIONS',14,GRAY,True)
text(s,.72,2.68,5.6,.91,'81.69%',59,BLUE,True)
text(s,.75,3.65,5.7,.4,'75.64% guardrail only → +6.05 pp',21,INK,True)
text(s,7.0,2.15,5.5,.36,'HARMFUL ALLOWS',14,GRAY,True)
text(s,7.0,2.68,5.5,.91,'−87.44%',54,TEAL,True)
text(s,7.03,3.65,5.4,.4,'4,738 → 595 / 15,840 harmful requests',19,INK,True)
rect(s,.7,4.3,11.95,.035,LINE)
headers=['Full accounting','Guardrail only','Monitored']
for x,w,t in zip([.76,6.72,9.78],[5.55,2.3,2.5],headers):text(s,x,4.48,w,.3,t,14,GRAY,True)
for i,(label,av,bv) in enumerate([('Safe refusal rate','18.44%','22.47%'),('Held requests','0','1,685 / 5.49%'),('Runtime model-input evaluations','30,720','152,150')]):
    y=4.95+i*.43;text(s,.76,y,5.55,.31,label,17,INK);text(s,6.72,y,2.3,.31,av,17,INK);text(s,9.78,y,2.5,.31,bv,17,BLUE,True)
rect(s,.7,6.43,11.95,.4,LIGHT)
text(s,.86,6.5,11.6,.25,'55/55 incidents confirmed · 45/50 fault episodes restored · 5 held · 0/10 false control confirmations',14,BLUE,True)
footer(s,8,'Evidence: benchmark-report.json + independent-audit.json · Holds receive zero automatic-correctness credit; not production-rate estimates')

# 09: verification gate, not unbounded self-repair.
s=base(9,'Reliability + responsible AI','Verification is the activation gate.','The actual candidate must pass—not just its plan or cached output.')
steps=[('01','Execute','Run the candidate on\nseparate verification probes.'),('02','Verify','Check both score vectors\nand decision outcomes.'),('03','Reverify','Check again before a\ngeneration-bound switch.')]
for x,(num,title,body) in zip([.75,4.95,9.15],steps):
    text(s,x,2.25,3.4,.35,num,16,BLUE,True,font='Courier New');rect(s,x,2.8,3.45,.04,BLUE)
    text(s,x,3.14,3.45,.45,title,29,INK,True);text(s,x,3.92,3.45,1.2,body,20,GRAY)
rect(s,.72,5.5,7.33,1.02,LIGHT,None,True)
text(s,.95,5.69,6.85,.64,'Explicit approval + single-use plan\nWeights and policy thresholds stay fixed.',19,BLUE,True)
rect(s,8.32,5.5,4.28,1.02,'D9FBFB',None,True)
text(s,8.55,5.69,3.82,.64,'No verified replacement?\nKeep an explicit hold.',19,TEAL,True)
text(s,.77,6.69,11.85,.23,'Runtime restoration preserves baseline behavior. Semantic classifier quality still requires labelled evaluation.',12.5,GRAY)
footer(s,9,'Evidence: RecoveryManager.propose/apply · verify_outputs · SupervisedRuntime.release_hold · Full scope and historical recalibration in notes')

# 10: LinuxONE is the implementation foundation.
s=base(10,'IBM Z integration + conclusion','Native IBM LinuxONE. End to end.','A complete local CPU workflow: generate → fit → infer → detect → verify → audit.')
text(s,.75,2.3,6.0,1.6,'s390x. Big-endian Linux.\nNo external model API.\nFresh offline reproduction.',27,INK,True)
text(s,.78,4.27,5.65,1.1,'Qualified container:\nTwo-CPU quota · 4 GiB memory limit',20,GRAY)
for x,y,value,label in [(7.1,2.2,'39.82 s','Selected classifier fit'),(10.05,2.2,'923.84 MiB','Peak child-process RSS'),(7.1,4.05,'292','Native tests pass'),(10.05,4.05,'100,000','Data rows reproduced')]:
    rect(s,x,y,2.55,.035,BLUE);text(s,x,y+.25,2.6,.65,value,30 if len(value)<8 else 24,BLUE,True);text(s,x,y+.98,2.6,.65,label,16,GRAY)
rect(s,.72,5.94,11.9,.72,NAVY,None,True)
text(s,.97,6.1,11.4,.4,'A guardrail needs a watcher. Every recovery needs evidence.',25,WHITE,True)
footer(s,10,'Native CPU qualification · 4 optional PyArrow tests skip · No Telum/Spyre, attestation, hardware-speedup, or production-capacity claim')

prs.save(OUT)
# Deterministic structural checks; numerical claims come from the frozen source tables.
assert len(prs.slides)==10
for i,slide in enumerate(prs.slides,1):
    assert slide.notes_slide.notes_text_frame.text.strip()
    assert 'SOURCE REFERENCE INDEX' in slide.notes_slide.notes_text_frame.text
    for shape in slide.shapes:
        if shape.width>0 and shape.height>0:
            assert shape.left>=0 and shape.top>=0,(i,'negative geometry')
            assert shape.left+shape.width<=prs.slide_width+20,(i,'horizontal overflow')
            assert shape.top+shape.height<=prs.slide_height+20,(i,'vertical overflow')
print(json.dumps({'powerpoint':str(OUT),'slides':10,'bytes':OUT.stat().st_size,'editable_text':True,'speaker_notes_with_sources':True},indent=2))
